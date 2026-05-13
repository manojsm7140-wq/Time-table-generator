"""
Timetable Views - Dashboard, upload, generate, view timetable
"""
import json
from collections import defaultdict
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.http import JsonResponse, HttpResponse
from django.views.decorators.http import require_POST

from management.models import ExcelUpload, Subject, Room, Class, StaffProfile, Department
from .models import TimetableEntry, ConflictLog
from .excel_parser import (
    parse_student_timetable, parse_room_mapping,
    detect_conflicts, build_timetable_grid, DAYS_ORDERED, TIME_SLOT_MAP
)

TIME_SLOTS_DISPLAY = [
    ('08:00-09:00', '8:00-9:00'),
    ('09:00-10:00', '9:00-10:00'),
    ('10:00-11:00', '10:00-11:00'),
    ('11:00-12:00', '11:00-12:00'),
    ('12:00-13:00', '12:00-1:00'),
    ('13:00-14:00', '1:00-2:00'),
    ('14:30-15:30', '2:30-3:30'),
    ('15:30-16:30', '3:30-4:30'),
    ('16:30-17:30', '4:30-5:30'),
]

DAYS_FULL = {
    'MON': 'Monday', 'TUE': 'Tuesday', 'WED': 'Wednesday',
    'THU': 'Thursday', 'FRI': 'Friday', 'SAT': 'Saturday',
}


@login_required
def dashboard(request):
    user = request.user
    ctx = {
        'total_classes': TimetableEntry.objects.values('class_label').distinct().count(),
        'total_conflicts': ConflictLog.objects.filter(resolved=False).count(),
        'total_entries': TimetableEntry.objects.count(),
        'recent_uploads': ExcelUpload.objects.order_by('-uploaded_at')[:5],
        'class_labels': TimetableEntry.objects.values_list('class_label', flat=True).distinct().order_by('class_label'),
        'conflicts': ConflictLog.objects.filter(resolved=False).order_by('-detected_at')[:5],
    }
    if user.role == 'staff':
        try:
            profile = user.staffprofile
            ctx['staff_entries'] = TimetableEntry.objects.filter(staff=profile).order_by('day', 'time_slot')
        except Exception:
            ctx['staff_entries'] = []
    return render(request, 'timetable/dashboard.html', ctx)


@login_required
def upload_excel(request):
    if request.user.role not in ('admin',):
        messages.error(request, 'Only admins can upload files.')
        return redirect('timetable:dashboard')

    if request.method == 'POST':
        upload_type = request.POST.get('upload_type')
        file = request.FILES.get('file')
        if not file:
            messages.error(request, 'No file selected.')
            return redirect('timetable:upload')

        upload = ExcelUpload.objects.create(
            uploaded_by=request.user,
            upload_type=upload_type,
            file=file
        )
        messages.success(request, f'File "{file.name}" uploaded successfully. Click "Process" to generate timetable.')
        return redirect('timetable:upload')

    uploads = ExcelUpload.objects.order_by('-uploaded_at')
    return render(request, 'timetable/upload.html', {'uploads': uploads})


@login_required
def process_upload(request, upload_id):
    if request.user.role != 'admin':
        messages.error(request, 'Permission denied.')
        return redirect('timetable:dashboard')

    upload = get_object_or_404(ExcelUpload, id=upload_id)
    filepath = upload.file.path
    all_errors = []
    entries_created = 0

    if upload.upload_type == 'timetable':
        entries, errors = parse_student_timetable(filepath)
        all_errors.extend(errors)

        # Clear existing entries
        TimetableEntry.objects.all().delete()
        ConflictLog.objects.all().delete()

        # Bulk create entries
        objs = []
        for e in entries:
            subj = Subject.objects.filter(code__iexact=e['subject_code']).first()
            room = Room.objects.filter(room_number=e['room_number']).first()
            objs.append(TimetableEntry(
                class_label=e['class_label'],
                day=e['day'],
                time_slot=e['time_slot'],
                raw_entry=e['raw_entry'],
                raw_subject_code=e['subject_code'],
                raw_room_number=e['room_number'],
                subject=subj,
                room=room,
            ))
        TimetableEntry.objects.bulk_create(objs, ignore_conflicts=True)
        entries_created = len(objs)

        # Detect conflicts
        conflicts = detect_conflicts(entries)
        for c in conflicts:
            ConflictLog.objects.create(**c)

    elif upload.upload_type == 'room_mapping':
        room_data, errors = parse_room_mapping(filepath)
        all_errors.extend(errors)
        # Auto-create rooms from mapping
        rooms_created = 0
        for day, slots in room_data.items():
            for slot in slots:
                rn = slot['room_number']
                if not Room.objects.filter(room_number=rn).exists():
                    Room.objects.create(room_number=rn, room_type='classroom')
                    rooms_created += 1
        entries_created = rooms_created
        messages.success(request, f'Room mapping processed. {rooms_created} new rooms created.')

    upload.processed = True
    upload.notes = f"Processed: {entries_created} entries. Errors: {len(all_errors)}"
    upload.save()

    if all_errors:
        messages.warning(request, f'Processed with {len(all_errors)} warnings: ' + '; '.join(all_errors[:3]))
    else:
        messages.success(request, f'Successfully processed! {entries_created} timetable entries loaded.')

    return redirect('timetable:view_timetable')


@login_required
def view_timetable(request):
    class_labels = list(TimetableEntry.objects.values_list('class_label', flat=True).distinct().order_by('class_label'))
    selected_class = request.GET.get('class_label', class_labels[0] if class_labels else '')
    view_type = request.GET.get('view', 'class')

    grid = {day: {slot[0]: [] for slot in TIME_SLOTS_DISPLAY} for day in DAYS_ORDERED}
    entries = []

    if view_type == 'class' and selected_class:
        entries = TimetableEntry.objects.filter(class_label=selected_class).order_by('day', 'time_slot')
        for entry in entries:
            day = entry.day
            slot = entry.time_slot
            if day in grid and slot in grid[day]:
                grid[day][slot].append(entry)

    elif view_type == 'room':
        selected_room = request.GET.get('room_number', '')
        rooms = Room.objects.order_by('room_number')
        if selected_room:
            entries = TimetableEntry.objects.filter(raw_room_number=selected_room).order_by('day', 'time_slot')
            for entry in entries:
                if entry.day in grid and entry.time_slot in grid[entry.day]:
                    grid[entry.day][entry.time_slot].append(entry)
        return render(request, 'timetable/view.html', {
            'grid': grid, 'view_type': view_type,
            'rooms': rooms, 'selected_room': selected_room,
            'time_slots': TIME_SLOTS_DISPLAY, 'days': DAYS_ORDERED,
            'days_full': DAYS_FULL, 'class_labels': class_labels,
        })

    conflicts = ConflictLog.objects.filter(resolved=False)

    ctx = {
        'grid': grid,
        'view_type': view_type,
        'class_labels': class_labels,
        'selected_class': selected_class,
        'time_slots': TIME_SLOTS_DISPLAY,
        'days': DAYS_ORDERED,
        'days_full': DAYS_FULL,
        'conflicts': conflicts,
        'conflict_count': conflicts.count(),
        'entries': entries,
    }
    return render(request, 'timetable/view.html', ctx)


@login_required
def staff_timetable(request):
    """View for staff to see their own schedule."""
    if request.user.role not in ('staff', 'admin'):
        return redirect('timetable:dashboard')
    
    staff_id = request.GET.get('staff_id')
    if request.user.role == 'admin' and staff_id:
        profile = get_object_or_404(StaffProfile, id=staff_id)
    else:
        try:
            profile = request.user.staffprofile
        except Exception:
            messages.info(request, 'No staff profile found for your account.')
            return redirect('timetable:dashboard')
    
    entries = TimetableEntry.objects.filter(staff=profile).order_by('day', 'time_slot')
    grid = {day: {slot[0]: [] for slot in TIME_SLOTS_DISPLAY} for day in DAYS_ORDERED}
    for entry in entries:
        if entry.day in grid and entry.time_slot in grid[entry.day]:
            grid[entry.day][entry.time_slot].append(entry)
    
    all_staff = StaffProfile.objects.select_related('user').all() if request.user.role == 'admin' else []
    
    return render(request, 'timetable/staff_view.html', {
        'profile': profile, 'grid': grid, 'entries': entries,
        'time_slots': TIME_SLOTS_DISPLAY, 'days': DAYS_ORDERED,
        'days_full': DAYS_FULL, 'all_staff': all_staff,
    })


@login_required
def conflict_report(request):
    conflicts = ConflictLog.objects.order_by('-detected_at')
    if request.method == 'POST':
        conflict_id = request.POST.get('conflict_id')
        if conflict_id:
            ConflictLog.objects.filter(id=conflict_id).update(resolved=True)
            messages.success(request, 'Conflict marked as resolved.')
    return render(request, 'timetable/conflicts.html', {'conflicts': conflicts})


@login_required
def regenerate(request):
    """Re-process the last uploaded timetable file."""
    if request.user.role != 'admin':
        messages.error(request, 'Permission denied.')
        return redirect('timetable:dashboard')

    last_upload = ExcelUpload.objects.filter(upload_type='timetable').order_by('-uploaded_at').first()
    if last_upload:
        return process_upload(request, last_upload.id)
    messages.warning(request, 'No timetable file uploaded yet.')
    return redirect('timetable:upload')


# ─── EXPORT VIEWS ────────────────────────────────────────────
@login_required
def export_excel(request):
    """Export timetable for a class section as .xlsx"""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from django.http import HttpResponse

    class_label = request.GET.get('class_label', '')
    if not class_label:
        messages.error(request, 'Please select a class first.')
        return redirect('timetable:view_timetable')

    entries = TimetableEntry.objects.filter(class_label=class_label).order_by('day', 'time_slot')

    wb = openpyxl.Workbook()
    ws = wb.active
    import re as _re
    safe_title = _re.sub(r'[\\/*?:\[\]]', '-', class_label)[:31]
    ws.title = safe_title

    # Styles
    header_font = Font(bold=True, color='FFFFFF', size=11)
    header_fill = PatternFill('solid', fgColor='1a3a6b')
    day_font = Font(bold=True, color='FFFFFF', size=10)
    day_fill = PatternFill('solid', fgColor='2d5a9e')
    cell_font = Font(size=9)
    center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    thin = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )

    # Title row
    ws.merge_cells('A1:H1')
    ws['A1'] = f'JSS College Mysore - {class_label} - Even Semester 2025-26'
    ws['A1'].font = Font(bold=True, size=13, color='1a3a6b')
    ws['A1'].alignment = center
    ws.row_dimensions[1].height = 22

    # Header row
    headers = ['Time Slot', 'Monday', 'Tuesday', 'Wednesday', 'Thursday', 'Friday', 'Saturday']
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=2, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = thin
    ws.row_dimensions[2].height = 18

    # Build grid
    from collections import defaultdict
    grid = defaultdict(dict)
    for entry in entries:
        grid[entry.time_slot][entry.day] = entry.raw_entry

    row = 3
    for slot, slot_label in TIME_SLOTS_DISPLAY:
        ws.cell(row=row, column=1, value=slot_label).font = Font(bold=True, size=9)
        ws.cell(row=row, column=1).alignment = center
        ws.cell(row=row, column=1).border = thin
        ws.cell(row=row, column=1).fill = PatternFill('solid', fgColor='eef2ff')

        for col_idx, day in enumerate(DAYS_ORDERED, 2):
            val = grid[slot].get(day, '')
            cell = ws.cell(row=row, column=col_idx, value=val)
            cell.font = cell_font
            cell.alignment = center
            cell.border = thin
            if val:
                cell.fill = PatternFill('solid', fgColor='dbeafe')
        ws.row_dimensions[row].height = 28
        row += 1

    # Column widths
    ws.column_dimensions['A'].width = 14
    for col in 'BCDEFGH':
        ws.column_dimensions[col].width = 18

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    safe_name = class_label.replace(' ', '_').replace('/', '-')[:40]
    response['Content-Disposition'] = f'attachment; filename="timetable_{safe_name}.xlsx"'
    wb.save(response)
    return response


@login_required
def export_pdf(request):
    """Export timetable as PDF using ReportLab."""
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm
    from django.http import HttpResponse
    from io import BytesIO
    from collections import defaultdict

    class_label = request.GET.get('class_label', '')
    if not class_label:
        messages.error(request, 'Please select a class first.')
        return redirect('timetable:view_timetable')

    entries = TimetableEntry.objects.filter(class_label=class_label).order_by('day', 'time_slot')

    grid = defaultdict(dict)
    for entry in entries:
        grid[entry.time_slot][entry.day] = entry.raw_entry

    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(A4),
                            rightMargin=1*cm, leftMargin=1*cm,
                            topMargin=1.5*cm, bottomMargin=1*cm)

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle('title', parent=styles['Heading1'],
                                  fontSize=13, textColor=colors.HexColor('#1a3a6b'),
                                  spaceAfter=8)
    sub_style = ParagraphStyle('sub', parent=styles['Normal'],
                                fontSize=9, textColor=colors.grey, spaceAfter=12)

    elements = [
        Paragraph('JSS College, Ooty Road, Mysore – 570025', title_style),
        Paragraph(f'{class_label} | Even Semester 2025-26', sub_style),
    ]

    # Table data
    header_row = ['Time'] + [DAYS_FULL[d] for d in DAYS_ORDERED]
    data = [header_row]
    for slot, slot_label in TIME_SLOTS_DISPLAY:
        row_data = [slot_label]
        for day in DAYS_ORDERED:
            row_data.append(grid[slot].get(day, ''))
        data.append(row_data)

    col_widths = [3.2*cm] + [3.6*cm]*6
    table = Table(data, colWidths=col_widths, repeatRows=1)
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1a3a6b')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 9),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('BACKGROUND', (0, 1), (0, -1), colors.HexColor('#eef2ff')),
        ('FONTNAME', (0, 1), (0, -1), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 1), (-1, -1), 8),
        ('ROWBACKGROUNDS', (1, 1), (-1, -1), [colors.white, colors.HexColor('#f8faff')]),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#c7d2fe')),
        ('ROWHEIGHT', (0, 0), (-1, -1), 28),
    ]))

    elements.append(table)
    doc.build(elements)
    buffer.seek(0)

    response = HttpResponse(buffer, content_type='application/pdf')
    safe_name = class_label.replace(' ', '_').replace('/', '-')[:40]
    response['Content-Disposition'] = f'attachment; filename="timetable_{safe_name}.pdf"'
    return response
